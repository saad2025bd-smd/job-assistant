"""Extract text and simple signals from uploaded CV/resume files."""
from typing import Dict, Any, List, Tuple
import re


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes using pypdf."""
    try:
        from pypdf import PdfReader
        from io import BytesIO
        reader = PdfReader(BytesIO(file_bytes))
        parts = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                parts.append(t)
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[PDF extract error: {e}]"


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        from docx import Document
        from io import BytesIO
        doc = Document(BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[DOCX extract error: {e}]"


def extract_text_from_upload(uploaded_file) -> Tuple[str, str]:
    """
    Returns (text, file_type_note).
    Supports PDF, DOCX, TXT.
    """
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        return extract_text_from_pdf(data), "pdf"
    if name.endswith(".docx"):
        return extract_text_from_docx(data), "docx"
    if name.endswith(".txt") or name.endswith(".md"):
        try:
            return data.decode("utf-8", errors="replace").strip(), "text"
        except Exception:
            return data.decode("latin-1", errors="replace").strip(), "text"
    # Fallback: try pdf then text
    text = extract_text_from_pdf(data)
    if text and not text.startswith("[PDF extract error"):
        return text, "pdf"
    try:
        return data.decode("utf-8", errors="replace").strip(), "text"
    except Exception:
        return "", "unknown"


# Common skill-like tokens to pull from CV text
SKILL_CANDIDATES = [
    # languages
    "python", "javascript", "typescript", "java", "c++", "c#", "go", "rust", "ruby",
    "php", "swift", "kotlin", "scala", "r", "matlab", "sql", "html", "css", "bash", "shell",
    # frameworks / libs
    "react", "angular", "vue", "next.js", "node.js", "django", "flask", "fastapi",
    "spring", "express", "nestjs", "rails", "laravel", "dotnet", ".net",
    "tensorflow", "pytorch", "keras", "scikit-learn", "pandas", "numpy",
    # cloud / devops
    "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "k8s", "terraform",
    "jenkins", "github actions", "gitlab ci", "ci/cd", "linux", "unix",
    # data / db
    "postgresql", "mysql", "mongodb", "redis", "elasticsearch", "snowflake", "bigquery",
    "spark", "hadoop", "airflow", "dbt",
    # tools / other
    "git", "jira", "figma", "salesforce", "tableau", "power bi", "excel",
    "rest", "graphql", "grpc", "microservices", "agile", "scrum",
    "machine learning", "deep learning", "nlp", "computer vision", "llm", "openai",
]


def extract_skills_from_text(text: str) -> List[str]:
    """Simple case-insensitive skill extraction from CV text."""
    lower = text.lower()
    found = []
    for skill in SKILL_CANDIDATES:
        # word-boundary-ish match
        pattern = r"(?<![a-z])" + re.escape(skill.lower()) + r"(?![a-z])"
        if re.search(pattern, lower):
            # Preserve a nice display form
            display = skill.upper() if skill in ("aws", "gcp", "sql", "html", "css", "nlp", "llm", "php", "ci/cd") else skill.title()
            if skill == "node.js":
                display = "Node.js"
            elif skill == "next.js":
                display = "Next.js"
            elif skill == "c++":
                display = "C++"
            elif skill == "c#":
                display = "C#"
            elif skill == ".net":
                display = ".NET"
            found.append(display)
    # dedupe preserving order
    seen = set()
    out = []
    for s in found:
        key = s.lower()
        if key not in seen:
            seen.add(key)
            out.append(s)
    return out


def summarize_cv_for_profile(text: str, max_chars: int = 1200) -> str:
    """Return a shortened version suitable for experience_summary."""
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rsplit(" ", 1)[0] + "…"


def analyze_cv(text: str) -> Dict[str, Any]:
    """Return structured signals extracted from CV text."""
    skills = extract_skills_from_text(text)
    summary = summarize_cv_for_profile(text)
    return {
        "skills_found": skills,
        "summary": summary,
        "char_count": len(text),
        "word_count": len(text.split()),
    }
